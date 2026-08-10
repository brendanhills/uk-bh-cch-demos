import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    
    # Color Palette
    COLOR_BG = RGBColor(255, 255, 255)         # Pure White
    COLOR_CARD = RGBColor(248, 249, 250)       # Light Grey Card fill (#F8F9FA)
    COLOR_BORDER = RGBColor(218, 220, 224)     # Border Grey (#DADCE0)
    
    # Google Brand Colors
    COLOR_BLUE = RGBColor(26, 115, 232)        # Google Blue (#1A73E8)
    COLOR_GREEN = RGBColor(52, 168, 83)        # Google Green (#34A853)
    COLOR_YELLOW = RGBColor(251, 188, 5)       # Google Yellow/Orange (#FBBC05)
    COLOR_RED = RGBColor(234, 67, 53)          # Google Red (#EA4335)
    
    # Text colors
    COLOR_TEXT_MAIN = RGBColor(32, 33, 36)     # Off-Black (#202124)
    COLOR_TEXT_MUTED = RGBColor(95, 99, 104)   # Medium Grey (#5F6368)

    # =========================================================================
    # SLIDE 1: PROBLEM FRAMING & SCENARIO
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    
    # Background
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_BG
    bg1.line.fill.background()
    
    # Top Header Box
    h_box1 = slide1.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.133), Inches(0.9))
    tf_h1 = h_box1.text_frame
    tf_h1.word_wrap = True
    tf_h1.margin_left = tf_h1.margin_right = tf_h1.margin_top = tf_h1.margin_bottom = 0
    
    p_h1 = tf_h1.paragraphs[0]
    r_h1 = p_h1.add_run()
    r_h1.text = "Elevating Pediatric Home Care"
    r_h1.font.name = 'Arial'
    r_h1.font.size = Pt(22)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_BLUE
    p_h1.space_after = Pt(2)
    
    p_sub1 = tf_h1.add_paragraph()
    r_s1 = p_sub1.add_run()
    r_s1.text = "Moving from theory to practice for Cymbal Children's Hospital Home Care Concierge"
    r_s1.font.name = 'Arial'
    r_s1.font.size = Pt(11)
    r_s1.font.color.rgb = COLOR_TEXT_MUTED
    
    # Left Content Area (DEMO 1 Badge + Large Headline)
    left_box1 = slide1.shapes.add_textbox(Inches(0.6), Inches(1.8), Inches(6.0), Inches(4.5))
    tf_l1 = left_box1.text_frame
    tf_l1.word_wrap = True
    tf_l1.margin_left = tf_l1.margin_right = tf_l1.margin_top = tf_l1.margin_bottom = 0
    
    # DEMO 1 Pill Badge
    p_badge = tf_l1.paragraphs[0]
    r_badge = p_badge.add_run()
    r_badge.text = "DEMO 1"
    r_badge.font.name = 'Arial'
    r_badge.font.size = Pt(10)
    r_badge.font.bold = True
    r_badge.font.color.rgb = COLOR_BLUE
    p_badge.space_after = Pt(16)
    
    # Main Headline
    p_main1 = tf_l1.add_paragraph()
    r_m1 = p_main1.add_run()
    r_m1.text = "Bridging the gap between "
    r_m1.font.name = 'Arial'
    r_m1.font.size = Pt(28)
    r_m1.font.bold = True
    r_m1.font.color.rgb = COLOR_TEXT_MAIN
    
    r_m2 = p_main1.add_run()
    r_m2.text = "hospital discharge "
    r_m2.font.name = 'Arial'
    r_m2.font.size = Pt(28)
    r_m2.font.bold = True
    r_m2.font.color.rgb = COLOR_BLUE
    
    r_m3 = p_main1.add_run()
    r_m3.text = "and home recovery"
    r_m3.font.name = 'Arial'
    r_m3.font.size = Pt(28)
    r_m3.font.bold = True
    r_m3.font.color.rgb = COLOR_TEXT_MAIN
    
    # Right Scenario Card Box
    scen_card = slide1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(1.8), Inches(5.7), Inches(4.5)
    )
    scen_card.fill.solid()
    scen_card.fill.fore_color.rgb = COLOR_CARD
    scen_card.line.color.rgb = COLOR_BORDER
    scen_card.line.width = Pt(1)
    
    scen_box = slide1.shapes.add_textbox(Inches(7.3), Inches(2.1), Inches(5.1), Inches(3.9))
    tf_scen = scen_box.text_frame
    tf_scen.word_wrap = True
    tf_scen.margin_left = tf_scen.margin_right = tf_scen.margin_top = tf_scen.margin_bottom = 0
    
    # SCENARIO Header
    p_sh = tf_scen.paragraphs[0]
    r_sh = p_sh.add_run()
    r_sh.text = "💬  SCENARIO"
    r_sh.font.name = 'Arial'
    r_sh.font.size = Pt(11)
    r_sh.font.bold = True
    r_sh.font.color.rgb = COLOR_BLUE
    p_sh.space_after = Pt(16)
    
    # Scenario Body Text (High-level problem framing without spoiling demo details)
    p_sb = tf_scen.add_paragraph()
    r_sb = p_sb.add_run()
    r_sb.text = "A family is managing a child's home recovery from Asthma/RSV, facing complex discharge instructions, specialized nurse scheduling friction, and financial confusion during high-stress recovery."
    r_sb.font.name = 'Arial'
    r_sb.font.size = Pt(13.5)
    r_sb.font.color.rgb = COLOR_TEXT_MAIN
    p_sb.space_after = Pt(24)
    
    # Highlight Question
    p_sq = tf_scen.add_paragraph()
    r_sq = p_sq.add_run()
    r_sq.text = "How can Gemini Live native audio and agentic workflows transform post-discharge recovery into seamless, connected care?"
    r_sq.font.name = 'Arial'
    r_sq.font.size = Pt(13.5)
    r_sq.font.bold = True
    r_sq.font.color.rgb = COLOR_BLUE

    # =========================================================================
    # SLIDE 2: PIPELINE & WORKFLOW
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    
    # Background
    bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = COLOR_BG
    bg2.line.fill.background()
    
    # Header Text Box
    h_box2 = slide2.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.133), Inches(0.9))
    tf_h2 = h_box2.text_frame
    tf_h2.word_wrap = True
    tf_h2.margin_left = tf_h2.margin_right = tf_h2.margin_top = tf_h2.margin_bottom = 0
    
    p_h2 = tf_h2.paragraphs[0]
    r_h2 = p_h2.add_run()
    r_h2.text = "Clinical Voice & Care Coordination Pipeline"
    r_h2.font.name = 'Arial'
    r_h2.font.size = Pt(22)
    r_h2.font.bold = True
    r_h2.font.color.rgb = COLOR_BLUE
    p_h2.space_after = Pt(2)
    
    p_sub2 = tf_h2.add_paragraph()
    r_s2 = p_sub2.add_run()
    r_s2.text = "Real-time consultation, automated subsidy booking & hospital EMR documentation flow"
    r_s2.font.name = 'Arial'
    r_s2.font.size = Pt(11)
    r_s2.font.color.rgb = COLOR_TEXT_MUTED
    
    # 4 Columns Grid Dimensions
    col_w = Inches(2.8)
    col_h = Inches(5.3)
    col_y = Inches(1.5)
    col_gap = Inches(0.31)
    
    c1_x = Inches(0.6)
    c2_x = c1_x + col_w + col_gap
    c3_x = c2_x + col_w + col_gap
    c4_x = c3_x + col_w + col_gap
    
    p_x = Inches(0.14)
    p_y = Inches(0.15)
    
    pipeline_data = [
        {
            "x": c1_x,
            "step": "🎙️  01 / Hearing",
            "title": "Speech & Intake Capture",
            "code": "gemini-2.5-flash-native-audio",
            "accent": COLOR_BLUE,
            "bullets": [
                "Captures native speech-to-speech (<0.9s latency)",
                "Parses uploaded discharge summaries & PDFs",
                "Affective Dialog detects parental stress tone"
            ]
        },
        {
            "x": c2_x,
            "step": "🌐  02 / Simplifying",
            "title": "Jargon & Translation",
            "code": "gemini-live-api",
            "accent": COLOR_GREEN,
            "bullets": [
                "Multi-lingual support across global languages",
                "Simplifies complex spacer dosing for laypersons",
                "Grounded queries bypass AI hallucination risks"
            ]
        },
        {
            "x": c3_x,
            "step": "🔊  03 / Actioning",
            "title": "Care & Subsidy Booking",
            "code": "google-adk-toolkit",
            "accent": COLOR_YELLOW,
            "bullets": [
                "Calculates home visit costs & NDIS subsidies",
                "Books specialized home care nurse slots in-call",
                "Instant emergency 000 triage safety intercept"
            ]
        },
        {
            "x": c4_x,
            "step": "📄  04 / Documenting",
            "title": "EMR SOAP Integration",
            "code": "cch-clinical-emr",
            "accent": COLOR_RED,
            "bullets": [
                "Syncs home vital logs & visit notes in real-time",
                "Commits approved funding records to central EMR",
                "Generates interactive patient recovery summaries"
            ]
        }
    ]
    
    # Generate 4 Columns
    for col in pipeline_data:
        # Card Background
        card = slide2.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, col["x"], col_y, col_w, col_h
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)
        
        # Top Accent Line
        accent_bar = slide2.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, col["x"], col_y, col_w, Inches(0.12)
        )
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = col["accent"]
        accent_bar.line.fill.background()
        
        # Text Box
        t_box = slide2.shapes.add_textbox(
            col["x"] + p_x, col_y + p_y + Inches(0.1), col_w - (p_x * 2), col_h - (p_y * 2) - Inches(0.1)
        )
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        
        # Step Sub-header
        p_step = tf.paragraphs[0]
        r_step = p_step.add_run()
        r_step.text = col["step"]
        r_step.font.name = 'Arial'
        r_step.font.size = Pt(10)
        r_step.font.bold = True
        r_step.font.color.rgb = col["accent"]
        p_step.space_after = Pt(2)
        
        # Main Pillar Title
        p_title = tf.add_paragraph()
        r_title = p_title.add_run()
        r_title.text = col["title"]
        r_title.font.name = 'Arial'
        r_title.font.size = Pt(13)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_TEXT_MAIN
        p_title.space_after = Pt(4)
        
        # Tech Code Pill Box
        p_code = tf.add_paragraph()
        r_code = p_code.add_run()
        r_code.text = col["code"]
        r_code.font.name = 'Courier New'
        r_code.font.size = Pt(8.5)
        r_code.font.bold = True
        r_code.font.color.rgb = col["accent"]
        p_code.space_after = Pt(12)
        
        # Bullets
        for b_text in col["bullets"]:
            p_b = tf.add_paragraph()
            r_b = p_b.add_run()
            r_b.text = "• " + b_text
            r_b.font.name = 'Arial'
            r_b.font.size = Pt(9.5)
            r_b.font.color.rgb = COLOR_TEXT_MUTED
            p_b.space_after = Pt(8)

    # Save Presentation
    from pathlib import Path
    output_path = Path(__file__).parent.parent / "docs" / "cch_problem_framing.pptx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    print(f"Successfully re-compiled presentation focused purely on high-level problem framing at {output_path}!")

if __name__ == "__main__":
    create_presentation()
